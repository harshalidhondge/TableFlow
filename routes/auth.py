import os
from flask import (
    Blueprint,
    render_template,
    request,
    redirect,
    url_for,
    session,
    flash,
    send_file,
    current_app
)

from datetime import date

from models.user import User
from models.table import Table
from reports import generate_pdf_report
from utils.extensions import db, bcrypt

from reportlab.platypus import (
    SimpleDocTemplate,
    Table as PDFTable,
    TableStyle
)

from reportlab.lib import colors

from openpyxl import Workbook
from openpyxl.styles import Font

import os

auth = Blueprint("auth", __name__)

ALLOWED_IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "gif"}


def is_allowed_image(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_IMAGE_EXTENSIONS

# ==========================================
# REGISTER
# ==========================================

@auth.route("/register", methods=["GET", "POST"])
def register():

    if request.method == "POST":

        name = request.form["name"]
        email = request.form["email"]
        phone = request.form["phone"]
        password = request.form["password"]

        existing = User.query.filter_by(
            email=email
        ).first()

        if existing:

            flash(
                "Email already registered.",
                "danger"
            )

            return redirect(
                url_for("auth.register")
            )

        hashed = bcrypt.generate_password_hash(
            password
        ).decode("utf-8")

        user = User(

            name=name,

            email=email,

            phone=phone,

            password=hashed

        )

        db.session.add(user)

        db.session.commit()

        flash(
            "Registration Successful. Please Login.",
            "success"
        )

        return redirect(
            url_for("auth.login")
        )

    return render_template(
        "auth/register.html"
    )

# ==========================================
# LOGIN
# ==========================================

@auth.route("/login", methods=["GET","POST"])
def login():

    if request.method=="POST":

        email=request.form["email"]

        password=request.form["password"]

        user=User.query.filter_by(
            email=email
        ).first()

        if user and bcrypt.check_password_hash(
            user.password,
            password
        ):

            session["user_id"]=user.id

            flash(
                f"Welcome {user.name}",
                "success"
            )

            return redirect(
                url_for("auth.dashboard")
            )

        flash(
            "Invalid Email or Password",
            "danger"
        )

    return render_template(
        "auth/login.html"
    )

# ==========================================
# OWNER DASHBOARD
# ==========================================

@auth.route("/dashboard")
def dashboard():

    if "user_id" not in session:

        return redirect(
            url_for("auth.login")
        )

    user = User.query.get(
        session["user_id"]
    )

    today = date.today()

    # ==========================
    # Restaurant Statistics
    # ==========================

    total_tables = Table.query.filter_by(
        user_id=user.id
    ).count()

    total_bookings = Table.query.filter_by(
        user_id=user.id
    ).count()

    available = Table.query.filter_by(
        user_id=user.id,
        status="Available"
    ).count()

    reserved = Table.query.filter_by(
        user_id=user.id,
        status="Reserved"
    ).count()

    booked = Table.query.filter_by(
        user_id=user.id,
        status="Booked"
    ).count()

    today = date.today()

    today_bookings = Table.query.filter_by(
        user_id=user.id,
        booking_date=today
    ).count()

    upcoming_bookings = Table.query.filter(
        Table.user_id == user.id,
        Table.booking_date > today
    ).order_by(
        Table.booking_date.asc()
    ).all()

    upcoming_bookings_count = len(upcoming_bookings)

    occupancy = 0

    if total_tables > 0:
        occupancy = round(
            ((reserved + booked) / total_tables) * 100,
            2
        )

    recent_tables = Table.query.filter_by(
        user_id=user.id
    ).order_by(
        Table.id.desc()
    ).limit(5).all()

    tables = Table.query.filter_by(
        user_id=user.id
    ).order_by(
        Table.id.asc()
    ).all()

    return render_template(

    "dashboard/dashboard.html",

    user=user,

    restaurant_status="OPEN",

    total_tables=total_tables,

    total_bookings=total_bookings,

    today_bookings=today_bookings,

    upcoming_bookings=upcoming_bookings,

    upcoming_bookings_count=upcoming_bookings_count,

    available=available,

    reserved=reserved,

    booked=booked,

    occupancy=occupancy,

    recent_tables=recent_tables,
    tables=tables

)
# ==========================================
# CREATE BOOKING
# ==========================================

@auth.route("/create_table", methods=["GET", "POST"])
def create_table():

    if "user_id" not in session:
        return redirect(url_for("auth.login"))

    if request.method == "POST":

        table_name = request.form["table_name"]
        customer_name = request.form["customer_name"]
        description = request.form["description"]
        booking_date = request.form["booking_date"]
        booking_time = request.form["booking_time"]
        floor = request.form["floor"]
        capacity = request.form["capacity"]
        location = request.form["location"]
        table_type = request.form["table_type"]
        status = request.form["status"]

        # Duplicate Booking Check
        duplicate = Table.query.filter_by(
            table_name=table_name,
            booking_date=booking_date,
            booking_time=booking_time
        ).first()

        if duplicate:

            flash(
                "This table is already booked for the selected date and time.",
                "danger"
            )

            return redirect(
                url_for("auth.create_table")
            )

        booking = Table(

            table_name=table_name,

            customer_name=customer_name,

            description=description,

            booking_date=booking_date,

            booking_time=booking_time,

            floor=floor,

            capacity=capacity,

            location=location,

            table_type=table_type,

            status=status,

            user_id=session["user_id"]

        )

        db.session.add(booking)
        db.session.commit()

        flash(
            "Booking Created Successfully!",
            "success"
        )

        return redirect(
            url_for("auth.view_tables")
        )

    return render_template(
        "customer/create_table.html"
    )


# ==========================================
# VIEW BOOKINGS
# ==========================================

@auth.route("/view_tables")
def view_tables():

    if "user_id" not in session:
        return redirect(url_for("auth.login"))

    search = request.args.get(
        "search",
        ""
    )

    status = request.args.get(
        "status",
        ""
    )

    booking_date = request.args.get(
        "booking_date",
        ""
    )

    query = Table.query.filter_by(
        user_id=session["user_id"]
    )

    # Customer Search
    if search:

        query = query.filter(
            Table.customer_name.ilike(
                f"%{search}%"
            )
        )

    # Status Filter
    if status:

        query = query.filter(
            Table.status == status
        )

    # Date Filter
    if booking_date:

        query = query.filter(
            Table.booking_date == booking_date
        )

    bookings = query.order_by(
        Table.booking_date.desc(),
        Table.booking_time.desc()
    ).all()

    total = len(bookings)

    available_count = sum(
        1 for b in bookings
        if b.status == "Available"
    )

    reserved_count = sum(
        1 for b in bookings
        if b.status == "Reserved"
    )

    booked_count = sum(
        1 for b in bookings
        if b.status == "Booked"
    )

    return render_template(

        "customer/view_tables.html",

        tables=bookings,

        total=total,

        available_count=available_count,

        reserved_count=reserved_count,

        booked_count=booked_count,

        search=search,

        status=status,

        booking_date=booking_date

    )
# ==========================================
# EDIT BOOKING
# ==========================================

@auth.route("/edit_table/<int:id>", methods=["GET", "POST"])
def edit_table(id):

    if "user_id" not in session:
        return redirect(url_for("auth.login"))

    booking = Table.query.get_or_404(id)

    if booking.user_id != session["user_id"]:

        flash(
            "Access Denied!",
            "danger"
        )

        return redirect(url_for("auth.view_tables"))

    if request.method == "POST":

        table_name = request.form["table_name"]
        customer_name = request.form["customer_name"]
        description = request.form["description"]
        booking_date = request.form["booking_date"]
        booking_time = request.form["booking_time"]
        floor = request.form["floor"]
        capacity = request.form["capacity"]
        location = request.form["location"]
        table_type = request.form["table_type"]
        status = request.form["status"]

        duplicate = Table.query.filter(
            Table.id != id,
            Table.table_name == table_name,
            Table.booking_date == booking_date,
            Table.booking_time == booking_time
        ).first()

        if duplicate:

            flash(
                "Another booking already exists for this table.",
                "danger"
            )

            return redirect(
                url_for("auth.edit_table", id=id)
            )

        booking.table_name = table_name
        booking.customer_name = customer_name
        booking.description = description
        booking.booking_date = booking_date
        booking.booking_time = booking_time
        booking.status = status

        db.session.commit()

        flash(
            "Booking Updated Successfully!",
            "success"
        )

        return redirect(
            url_for("auth.view_tables")
        )

    return render_template(
        "customer/edit_table.html",
        table=booking
    )


# ==========================================
# DELETE BOOKING
# ==========================================

@auth.route("/delete_table/<int:id>")
def delete_table(id):

    if "user_id" not in session:
        return redirect(url_for("auth.login"))

    booking = Table.query.get_or_404(id)

    if booking.user_id != session["user_id"]:

        flash(
            "Access Denied!",
            "danger"
        )

        return redirect(
            url_for("auth.view_tables")
        )

    db.session.delete(booking)
    db.session.commit()

    flash(
        "Booking Deleted Successfully!",
        "success"
    )

    return redirect(
        url_for("auth.view_tables"))


# ==========================================
# EXPORT PDF
# ==========================================

@auth.route('/export_pdf')
def export_pdf():
    if "user_id" not in session:
        return redirect(url_for("auth.login"))

    try:
        bookings = Table.query.filter_by(user_id=session["user_id"]).all()

        total_tables = Table.query.filter_by(user_id=session["user_id"]).count()
        reserved = Table.query.filter_by(user_id=session["user_id"], status="Reserved").count()
        booked = Table.query.filter_by(user_id=session["user_id"], status="Booked").count()

        occupancy_rate = 0
        if total_tables:
            occupancy_rate = round(((reserved + booked) / total_tables) * 100, 2)

        pdf_buffer = generate_pdf_report(bookings, occupancy_rate=occupancy_rate)

        return send_file(
            pdf_buffer,
            as_attachment=True,
            download_name='Bella_Vista_Report.pdf',
            mimetype='application/pdf'
        )
    except Exception as e:
        flash(f"Error generating PDF: {str(e)}", "danger")
        return redirect(url_for("auth.dashboard"))


# ==========================================
# EXPORT EXCEL
# ==========================================

@auth.route("/export_excel")
def export_excel():

    if "user_id" not in session:
        return redirect(url_for("auth.login"))

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Bella Vista Bookings"

    headers = [

        "ID",

        "Table",

        "Customer",

        "Description",

        "Date",

        "Time",

        "Status"

    ]

    for col, head in enumerate(headers, 1):

        cell = sheet.cell(row=1, column=col)

        cell.value = head

        cell.font = Font(bold=True)

    bookings = Table.query.filter_by(
        user_id=session["user_id"]
    ).all()

    row = 2

    for booking in bookings:

        sheet.cell(row=row,column=1).value = booking.id
        sheet.cell(row=row,column=2).value = booking.table_name
        sheet.cell(row=row,column=3).value = booking.customer_name
        sheet.cell(row=row,column=4).value = booking.description
        sheet.cell(row=row,column=5).value = str(booking.booking_date)
        sheet.cell(row=row,column=6).value = str(booking.booking_time)
        sheet.cell(row=row,column=7).value = booking.status

        row += 1

    excel_name = "BellaVistaBookings.xlsx"

    workbook.save(excel_name)

    return send_file(
        excel_name,
        as_attachment=True
    )
# ==========================================
# MANAGE RESTAURANT TABLES
# ==========================================

@auth.route("/manage_tables")
def manage_tables():

    if "user_id" not in session:
        return redirect(url_for("auth.login"))

    tables = Table.query.filter_by(
        user_id=session["user_id"]
    ).order_by(
        Table.table_name.asc()
    ).all()

    return render_template(
        "owner/manage_tables.html",
        tables=tables
    )


# ==========================================
# ADD RESTAURANT TABLE
# ==========================================

@auth.route("/add_table", methods=["GET", "POST"])
def add_table():

    if "user_id" not in session:
        return redirect(url_for("auth.login"))

    if request.method == "POST":

        table_name = request.form["table_name"]
        capacity = request.form["capacity"]
        location = request.form["location"]
        table_type = request.form["table_type"]

        duplicate = Table.query.filter_by(
            user_id=session["user_id"],
            table_name=table_name
        ).first()

        if duplicate:

            flash(
                "Table already exists.",
                "danger"
            )

            return redirect(
                url_for("auth.add_table")
            )

        table = Table(

            table_name=table_name,

            capacity=capacity,

            location=location,

            table_type=table_type,

            customer_name="",

            description="",

            booking_date=None,

            booking_time=None,

            status="Available",

            user_id=session["user_id"]

        )

        db.session.add(table)

        db.session.commit()

        flash(
            "Restaurant Table Added Successfully!",
            "success"
        )

        return redirect(
            url_for("auth.manage_tables")
        )

    return render_template(
        "owner/add_table.html"
    )


# ==========================================
# EDIT RESTAURANT TABLE
# ==========================================

@auth.route("/edit_restaurant_table/<int:id>", methods=["GET", "POST"])
def edit_restaurant_table(id):

    if "user_id" not in session:
        return redirect(url_for("auth.login"))

    table = Table.query.get_or_404(id)

    if request.method == "POST":

        table.table_name = request.form["table_name"]

        table.capacity = request.form["capacity"]

        table.location = request.form["location"]

        table.table_type = request.form["table_type"]

        db.session.commit()

        flash(
            "Restaurant Table Updated Successfully!",
            "success"
        )

        return redirect(
            url_for("auth.manage_tables")
        )

    return render_template(
        "owner/edit_table.html",
        table=table
    )


# ==========================================
# DELETE RESTAURANT TABLE
# ==========================================

@auth.route("/delete_restaurant_table/<int:id>")
def delete_restaurant_table(id):

    if "user_id" not in session:
        return redirect(url_for("auth.login"))

    table = Table.query.get_or_404(id)

    db.session.delete(table)

    db.session.commit()

    flash(
        "Restaurant Table Deleted Successfully!",
        "success"
    )

    return redirect(
        url_for("auth.manage_tables")
    )
# ==========================================
# RESTAURANT SUMMARY
# ==========================================

@auth.route("/restaurant_summary")
def restaurant_summary():

    if "user_id" not in session:
        return redirect(url_for("auth.login"))

    user_id = session["user_id"]

    total_tables = Table.query.filter_by(
        user_id=user_id
    ).count()

    available = Table.query.filter_by(
        user_id=user_id,
        status="Available"
    ).count()

    reserved = Table.query.filter_by(
        user_id=user_id,
        status="Reserved"
    ).count()

    booked = Table.query.filter_by(
        user_id=user_id,
        status="Booked"
    ).count()

    today = date.today()

    today_bookings = Table.query.filter_by(
        user_id=user_id,
        booking_date=today
    ).count()

    occupancy = 0

    if total_tables > 0:

        occupancy = round(

            ((reserved + booked) / total_tables) * 100,

            2

        )

    return {

        "restaurant": "Bella Vista",

        "total_tables": total_tables,

        "today_bookings": today_bookings,

        "available": available,

        "reserved": reserved,

        "occupied": booked,

        "occupancy": occupancy

    }


# ==========================================
# OWNER PROFILE
# ==========================================

@auth.route("/owner_profile", methods=["GET", "POST"])
def owner_profile():

    if "user_id" not in session:
        return redirect(url_for("auth.login"))

    user = User.query.get(session["user_id"])

    if request.method == "POST":
        if "profile_image" in request.files:
            image_file = request.files["profile_image"]
            if image_file and image_file.filename:
                if not is_allowed_image(image_file.filename):
                    flash("Please upload a valid image file (jpg, png, jpeg, gif).", "danger")
                    return redirect(url_for("auth.owner_profile"))

                filename = f"user_{user.id}_{image_file.filename}"
                upload_path = os.path.join(current_app.config["UPLOAD_FOLDER"], filename)
                image_file.save(upload_path)
                user.profile_image = filename
                db.session.commit()
                session["profile_image"] = filename
                flash("Profile photo updated successfully!", "success")
                return redirect(url_for("auth.owner_profile"))

        flash("No image file was selected.", "warning")
        return redirect(url_for("auth.owner_profile"))

    return render_template("owner/profile.html", user=user)


# ==========================================
# HELPER FUNCTION
# ==========================================

def get_dashboard_data(user_id):

    total = Table.query.filter_by(
        user_id=user_id
    ).count()

    available = Table.query.filter_by(
        user_id=user_id,
        status="Available"
    ).count()

    reserved = Table.query.filter_by(
        user_id=user_id,
        status="Reserved"
    ).count()

    booked = Table.query.filter_by(
        user_id=user_id,
        status="Booked"
    ).count()

    return {

        "total": total,

        "available": available,

        "reserved": reserved,

        "booked": booked

    }


# ==========================================
# LOGOUT
# ==========================================

@auth.route("/logout")
def logout():

    session.clear()

    flash(

        "Thank you for using Bella Vista.",

        "success"

    )

    return redirect(
        url_for("auth.login")
    )